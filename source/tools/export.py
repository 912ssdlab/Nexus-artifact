"""Export ONLY the results named in this invocation's manifest."""
import csv
import json
import math
from pathlib import Path
import shutil
from openpyxl import Workbook
from openpyxl.styles import Font

METRICS = ['mean_us', 'tm90_us', 'tm95_us', 'tm99_us', 'tm99_9_us', 'tm99_99_us',
           'p90_us', 'p95_us', 'p99_us', 'p99_9_us', 'p99_99_us',
           'hit_ratio', 'read_hit_ratio', 'write_hit_ratio', 'read_mean_us', 'write_mean_us',
           'flash_read_count', 'flush_count', 'io_events_per_request', 'io_amplification_4k_equiv',
           'trace_access_GiB', 'read_4k_equiv_GiB', 'write_4k_equiv_GiB', 'total_io_4k_equiv_GiB']

def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))

def numeric(v):
    if v is None: return None
    try:
        n = float(v)
        return n if math.isfinite(n) else None
    except (TypeError, ValueError): return v

def write_csv(path, rows, columns=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows and columns is None: return
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=columns or list(rows[0])); w.writeheader(); w.writerows(rows)

def sheet_table(sheet, name, rows):
    # Always leave exactly two blank rows between blocks.
    start = sheet.max_row + 3 if sheet.cell(1, 1).value is not None else 1
    sheet.cell(start, 1, name).font = Font(bold=True)
    if not rows: return
    keys = list(rows[0])
    for c, key in enumerate(keys, 1): sheet.cell(start + 1, c, key).font = Font(bold=True)
    for r, row in enumerate(rows, start + 2):
        for c, key in enumerate(keys, 1): sheet.cell(r, c, numeric(row[key]))

def export(out):
    out = Path(out)
    manifest = json.loads((out / 'run_manifest.json').read_text())
    traces = manifest['traces']; mechanisms = manifest['mechanisms']; sizes = manifest['cache_sizes_mib']
    root = out / 'paper_data'
    for folder in ('Motivation1', 'Motivation2', 'Motivation3', 'Main_Experiments/Nexus_Admission'): (root / folder).mkdir(parents=True, exist_ok=True)
    book = Workbook(); book.remove(book.active)
    info = book.create_sheet('Overview')
    info.append(['Mode', 'SMOKE - NOT PAPER RESULTS' if manifest['smoke'] else 'Full trace; no sampling'])
    info.append(['Source', 'Only this output directory; no historical/reference values'])
    info.append(['Latency unit', 'us; TM90/95/99 = mean of slowest 10/5/1%'])
    info.append(['Hit ratio', 'fraction 0-1'])
    info.append(['I/O amplification', '4KiB-equivalent events / 64B trace bytes; not physical NAND WAF'])
    info.append(['Optimal windows', 'Page 256K; Cacheline 512K global trace requests'])
    records = {}; missing = []; all_rows = []
    for t in traces:
        for size in sizes:
            for m in mechanisms:
                path = out / 'runs' / t['trace'] / f'{m}-{size}MB' / 'result.json'
                if not path.exists():
                    missing.append({'Trace': t['trace'], 'Cache_MiB': size, 'Mechanism': m}); continue
                r = json.loads(path.read_text())
                if r['signature']['trace_sha256'] != t['trace_sha256'] or r['metrics']['requests'] != t['requests']:
                    raise ValueError(f'Input/result mismatch: {path}')
                v = r['metrics']
                v.update(trace_access_GiB=t['requests']*64/2**30,
                         read_4k_equiv_GiB=v['flash_read_count']*4096/2**30,
                         write_4k_equiv_GiB=v['flush_count']*4096/2**30,
                         total_io_4k_equiv_GiB=(v['flash_read_count']+v['flush_count'])*4096/2**30)
                records[t['trace'], size, m] = r
                all_rows.append({'Workload': t['alias'], 'Trace': t['trace'], 'Cache_MiB': size, 'Mechanism': m,
                                 **{k: r['metrics'][k] for k in METRICS}})
    write_csv(root / 'Main_Experiments/all_metrics_long.csv', all_rows)
    write_csv(root / 'missing_results.csv', missing, ['Trace', 'Cache_MiB', 'Mechanism'])
    for size in sizes:
        for directory, subset in [('Motivation1', [t for t in traces if t['motivation']]), ('Main_Experiments', traces)]:
            ws = book.create_sheet(f'{directory}_{size}MB')
            norm_ws = book.create_sheet(f'{directory}_{size}MB_norm')
            for metric in METRICS:
                rows = []; normalized = []
                for t in subset:
                    row = {'Workload': t['alias']}
                    for m in mechanisms: row[m] = records.get((t['trace'], size, m), {}).get('metrics', {}).get(metric)
                    base = row.get('Page')
                    norm = {'Workload': t['alias'], **{m: row[m] / base if base and row[m] is not None else None for m in mechanisms}}
                    rows.append(row); normalized.append(norm)
                write_csv(root / directory / f'{size}MB_{metric}.csv', rows)
                write_csv(root / directory / f'{size}MB_{metric}_normalized_to_Page.csv', normalized)
                if directory == 'Main_Experiments':
                    write_csv(root / directory / 'selected9_six_mechanisms_raw' / f'{size}MB_{metric}.csv', rows)
                    write_csv(root / directory / 'selected9_six_mechanisms_page_normalized' / f'{size}MB_{metric}.csv', normalized)
                    four = [m for m in ('Page', 'Cacheline', 'SkyByte', 'Nexus') if m in mechanisms]
                    write_csv(root / directory / 'selected9_four_mechanisms_raw' / f'{size}MB_{metric}.csv',
                              [{k:r[k] for k in ['Workload',*four]} for r in rows])
                sheet_table(ws, metric, rows); sheet_table(norm_ws, metric + ' / ordinary Page', normalized)
    scope = []; basic = []; temporal_all = []; all_basic = []
    for t in traces:
        folder = out / 'features' / t['trace']
        if not (folder / 'complete.json').exists(): continue
        f = read_csv(folder / 'summary.csv')[0]
        b = {'Workload': t['alias'], 'Trace': t['trace'], 'Requests': int(f['accesses']),
             'Read_ratio': int(f['read_count']) / int(f['accesses']), 'Avg_interval_ns': float(f['mean_request_interval_raw']),
             'Footprint_GB': float(f['active_cacheline_ws_mib']) * 1024**2 / 1e9}
        for filename, key, column in [('hot_share', 'hot_cl_share', 'HCS'), ('read_hot_share', 'read_hot_cl_share', 'RHCS')]:
            path = folder / (filename + '.json')
            extra = json.loads(path.read_text()) if path.exists() else {}
            b[column] = extra.get(key)
            if filename == 'hot_share' and 'unique_cls' in extra:
                b['Footprint_GB'] = extra['unique_cls'] * 64 / 1e9
        all_basic.append(b)
        if not t['motivation']: continue
        basic.append(b)
        scope.append({'Workload': t['alias'], 'Trace': t['trace'], 'Window': 4096,
                      'Mean_active_CL': float(f['mean_active_cl']), 'Median_active_CL': float(f['median_active_cl']),
                      'P90_active_CL': float(f['p90_active_cl']), 'Small': float(f['p_active_small_le4']),
                      'Medium': float(f['p_active_medium_5_16']), 'Large': float(f['p_active_large_gt16']),
                      'Morans_I': numeric(f['moran_i_mean_standard']),
                      'Undefined_Moran_fraction': float(f['p_full_active'])})
        temporal = []
        for row in read_csv(folder / 'temporal.csv'):
            temporal.append({'Workload': t['alias'], 'Trace': t['trace'], 'Progress_percent': int(row['bin']) + .5,
                             'Observations': int(row['observations']), 'Mean_active_CL': float(row['mean_active_cl']),
                             'Mean_coverage': float(row['mean_active_cl']) / 64,
                             'Trace_mean_active_CL': float(f['mean_active_cl']),
                             'Trace_mean_coverage': float(f['mean_active_cl']) / 64})
        temporal_all.extend(temporal)
        write_csv(root / 'Motivation3' / f'{t["trace"]}_coverage_temporal.csv', temporal)
        sheet_table(book.create_sheet(f'Temporal_{t["alias"]}'), '100 equal-request bins', temporal)
    write_csv(root / 'Motivation2/selected6_scope_distribution_moran.csv', scope)
    write_csv(root / 'Motivation2/scope_moran_plot_ready.csv', scope)
    write_csv(root / 'Motivation2/selected6_workload_features.csv', basic)
    write_csv(root / 'Main_Experiments/selected9_workload_features.csv', all_basic)
    write_csv(root / 'Motivation3/selected6_coverage_temporal_long.csv', temporal_all)
    sheet_table(book.create_sheet('Motivation2_Scope_Moran'), 'Small/Medium/Large plus Moran secondary axis', scope)
    sheet_table(book.create_sheet('Workload_features'), 'Measured from input trace', all_basic)
    distributions = []; means = []; curves = []
    for (trace, size, m), r in records.items():
        if m != 'Nexus': continue
        run = out / r['run_directory']; dest = root / 'Main_Experiments/Nexus_Admission'
        for filename in ('granularity_timeline.csv', 'granularity_future_alignment.csv'):
            if not (run / filename).exists(): raise ValueError(f'Missing Nexus telemetry: {run / filename}')
            shutil.copy2(run / filename, dest / f'{trace}_{size}MB_{filename}')
        timeline = read_csv(run / 'granularity_timeline.csv')
        count = sum(int(x['admission_count']) for x in timeline)
        weighted = lambda k: sum(float(x[k]) * int(x['admission_count']) for x in timeline) / count if count else None
        means.append({'Workload': r['alias'], 'Cache_MiB': size, 'Admissions': count,
                      'Mean_logical_selected_CL': weighted('mean_logical_selected_cl'),
                      'Mean_selected_CL': weighted('mean_selected_cl'), 'Mean_requested_new_CL': weighted('mean_requested_new_cl'),
                      'Mean_backend_sectors': weighted('mean_backend_sectors')})
        counted = 0
        for g in (1, 2, 4, 8, 16, 32, 64):
            n = sum(int(x[f'count_{g}cl']) for x in timeline); counted += n
            distributions.append({'Workload': r['alias'], 'Cache_MiB': size, 'Selected_CL': g,
                                  'Admission_count': n, 'Fraction': n/count if count else None})
        if counted > count: raise ValueError('Granularity distribution exceeds admission count')
        distributions.append({'Workload': r['alias'], 'Cache_MiB': size, 'Selected_CL': 'Other',
                              'Admission_count': count-counted, 'Fraction': (count-counted)/count if count else None})
        for row in read_csv(run / 'granularity_future_alignment.csv'):
            curves.append({'Workload': r['alias'], 'Cache_MiB': size, **{k: numeric(v) for k, v in row.items()}})
    write_csv(root / 'Main_Experiments/Nexus_Admission/mean_granularity.csv', means)
    write_csv(root / 'Main_Experiments/Nexus_Admission/granularity_distribution.csv', distributions)
    write_csv(root / 'Main_Experiments/Nexus_Admission/temporal_alignment_long.csv', curves)
    sheet_table(book.create_sheet('Nexus_mean'), 'Admission-weighted means', means)
    sheet_table(book.create_sheet('Nexus_distribution'), 'Selected mask size distribution, NOT logical community size', distributions)
    sheet_table(book.create_sheet('Nexus_temporal'), 'Future A4096 versus admission size', curves)
    for ws in book:
        ws.freeze_panes = 'B3'
        ws.column_dimensions['A'].width = 27
        for c in 'BCDEFGHIJKLMNOPQRSTUVWXYZ': ws.column_dimensions[c].width = 21
    book.save(root / 'paper_results.xlsx')
    (root / 'completion.json').write_text(json.dumps({'smoke': manifest['smoke'], 'complete_jobs': len(records),
      'expected_jobs': manifest['expected_jobs'], 'all_simulations_complete': not missing,
      'eligible_as_full_selected9_reproduction': not manifest['smoke'] and not missing and len(traces)==9 and len(mechanisms)==6 and len(sizes)==2,
      'missing_cells': 'blank, never zero-filled', 'historical_results_used': False}, indent=2))

if __name__ == '__main__':
    import argparse
    p = argparse.ArgumentParser(); p.add_argument('output', type=Path); export(p.parse_args().output)
